# -*- coding: utf-8 -*-
"""
==================================================================
  יוצר תבנית ריקה לקובץ רשימת הלקוחות
==================================================================
הרץ אותי פעם אחת כדי לקבל קובץ אקסל מוכן (customers.xlsx)
עם העמודות הנכונות ושתי שורות דוגמה. אחר כך פתח אותו,
מחק את הדוגמאות, ומלא את הלקוחות האמיתיים שלך.

הרצה:   python create_customer_template.py
==================================================================
"""

import os

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

import config


def create_template():
    # אם כבר קיים קובץ לקוחות - לא דורסים אותו בטעות.
    if os.path.exists(config.CUSTOMERS_FILE):
        print(f"⚠️  הקובץ '{config.CUSTOMERS_FILE}' כבר קיים - לא יצרתי חדש.")
        print("   אם אתה רוצה תבנית נקייה, שנה/מחק את הקובץ הקיים קודם.")
        return

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "לקוחות"
    sheet.sheet_view.rightToLeft = True

    # שורת כותרות - חייבת להתאים לשמות שב-config.py
    # עמודת ה-ח.פ אופציונלית אך מומלצת מאוד (משפרת מאוד את הדיוק).
    headers = [
        config.CUSTOMERS_NAME_COLUMN,
        config.CUSTOMERS_ID_COLUMN,
        config.CUSTOMERS_COMPANY_ID_COLUMN,  # "ח.פ"
    ]
    fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    # שתי שורות דוגמה - מחק אותן ומלא את הלקוחות שלך.
    # העמודה השלישית (ח.פ) - מלא אם ידוע; אפשר גם להשאיר ריק.
    sheet.append(["ישראל ישראלי בע\"מ", "1001", "514111222"])
    sheet.append(["חברת אבן וסיד בע\"מ", "1002", "514333444"])

    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 15
    sheet.column_dimensions["C"].width = 15

    workbook.save(config.CUSTOMERS_FILE)
    print(f"✅ נוצרה תבנית: '{config.CUSTOMERS_FILE}'")
    print("   פתח אותה, מחק את שורות הדוגמה, ומלא את הלקוחות האמיתיים.")


if __name__ == "__main__":
    create_template()
