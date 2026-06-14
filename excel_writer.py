# -*- coding: utf-8 -*-
"""
==================================================================
  כתיבת קובץ האקסל (ה"מזכיר" של הכלי)
==================================================================
מקבל את כל השורות שעובדו ובונה מהן קובץ אקסל מסודר
עם כל העמודות שביקשת, כולל צביעה לפי סטטוס:
ירוק = מוכן לקבלה, צהוב = דורש בדיקה ידנית.
==================================================================
"""

from dataclasses import dataclass
from datetime import datetime
from typing import List, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

import config


# שם גיליון הלקוחות הגלוי (מקור לרשימה הנפתחת + חיפוש עם מסנן).
CUSTOMER_SHEET_TITLE = "לקוחות"
CUSTOMER_NAMES_RANGE = "CustomerNames"


@dataclass
class ResultRow:
    """שורה אחת בטבלת הפלט - שיק אחד מעובד."""
    detected_name: Optional[str]   # שם שזוהה על השיק
    matched_name: Optional[str]    # לקוח מותאם ברשימה
    confidence: int                # רמת ודאות (0-100)
    check_number: Optional[str]
    bank_name: Optional[str]
    branch_number: Optional[str]
    account_number: Optional[str]
    due_date: Optional[str]
    amount: Optional[str]
    status: str                    # מוכן / בדיקה ידנית / שיק שחזר
    maven_reference: Optional[str] # מספר האסמכתא מהשורה ב-Maven (88635 וכו')
    image_path: Optional[str]      # נתיב לצילום (לתיעוד)
    deposit_date: Optional[str] = None     # תאריך ההפקדה מדף הבנק = תאריך הקבלה
    scanned_amount: Optional[str] = None   # סכום שנסרק מהשיק (להשוואה בלבד!)
    company_id: Optional[str] = None       # ח.פ שזוהה על השיק
    evidence: Optional[str] = None         # על מה התבססה ההתאמה (שם/ח.פ/סכום)
    matched_customer_id: Optional[str] = None   # מספר הלקוח במערכת (מהרשימה)
    matched_company_id: Optional[str] = None    # ח.פ הלקוח מהרשימה (להשוואה)


# הכותרות של העמודות, בסדר שביקשת.
HEADERS = [
    "לאשר?",            # <<< אתה ממלא: כתוב "כן" בשורות שברצונך להוציא להן קבלה
    "שם לקוח שזוהה",
    "לקוח מותאם ברשימה",
    "מספר לקוח במערכת",  # המזהה של הלקוח כפי שמופיע ברשימה/Maven
    "רמת ודאות",
    "ח.פ שזוהה בשיק",    # ה-ח.פ שזוהה בצילום השיק
    "ח.פ במערכת",        # ה-ח.פ של הלקוח המותאם, מהרשימה - להשוואה
    "מספר שיק",
    "בנק",
    "סניף",
    "חשבון",
    "תאריך הפקדה",      # תאריך ההפקדה מדף הבנק - זהו תאריך הקבלה שתוצא
    "תאריך פירעון",
    "סכום",              # <<< הסכום מדף הבנק - זהו הסכום הקובע שיוצא בקבלה
    "סכום בשיק (סריקה)", # הסכום שנסרק מצילום השיק - להשוואה בלבד (עלול לטעות)
    "סטטוס",
    "בסיס ההתאמה",
    "אסמכתא Maven",
    "קובץ צילום",
]

# צבעים לסטטוס
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

STATUS_READY = "מוכן לקבלה"
STATUS_REVIEW = "דורש בדיקה ידנית"
STATUS_BOUNCED = "שיק שחזר - לא להוציא קבלה"


def write_results(rows: List[ResultRow], output_file: str = config.OUTPUT_FILE,
                  customers=None):
    """
    בונה את קובץ האקסל מתוך רשימת השורות.
    אם הועברה רשימת לקוחות (customers) - מוסיף רשימה נפתחת בעמודת
    "לקוח מותאם ברשימה", כדי לאפשר בחירת לקוח ידנית בשורות לבירור.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "שיקים לבדיקה"
    # הצגה מימין לשמאל (עברית)
    sheet.sheet_view.rightToLeft = True

    # שלב 1: כותבים את שורת הכותרות ומעצבים אותה.
    for col_index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # שלב 2: כותבים שורה לכל שיק, וצובעים לפי הסטטוס.
    for row_index, r in enumerate(rows, start=2):
        # עמודת "לאשר?": מציעים "כן" כברירת מחדל לשורות מוכנות לקבלה,
        # וריק לשורות שדורשות בדיקה / שיקים שחזרו. אתה יכול לשנות ידנית.
        approve_default = "כן" if r.status == STATUS_READY else ""
        values = [
            approve_default,
            r.detected_name or "",
            r.matched_name or "",
            r.matched_customer_id or "",
            f"{r.confidence}%" if r.confidence else "",
            r.company_id or "",
            r.matched_company_id or "",
            r.check_number or "",
            r.bank_name or "",
            r.branch_number or "",
            r.account_number or "",
            r.deposit_date or "",
            r.due_date or "",
            r.amount or "",
            r.scanned_amount or "",
            r.status,
            r.evidence or "",
            r.maven_reference or "",
            r.image_path or "",
        ]
        # צבע לפי הסטטוס: ירוק=מוכן, אדום=שיק שחזר, צהוב=בדיקה ידנית.
        if r.status == STATUS_READY:
            fill = GREEN_FILL
        elif r.status == STATUS_BOUNCED:
            fill = RED_FILL
        else:
            fill = YELLOW_FILL
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # שלב 3: מרחיבים את העמודות שיהיה נוח לקרוא.
    # סדר: לאשר, שם שזוהה, לקוח מותאם, מספר לקוח, ודאות, ח.פ בשיק, ח.פ במערכת,
    #      מספר שיק, בנק, סניף, חשבון, תאריך הפקדה, תאריך פירעון, סכום (בנק),
    #      סכום בשיק (סריקה), סטטוס, בסיס, אסמכתא, צילום
    widths = [8, 22, 22, 14, 9, 13, 13, 12, 12, 7, 13, 13, 13, 11, 15, 20, 24, 13, 22]
    for col_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[
            openpyxl.utils.get_column_letter(col_index)
        ].width = width

    # שלב 4: רשימה נפתחת לבחירת לקוח ידנית (אם יש רשימת לקוחות).
    if customers:
        _add_customer_dropdown(workbook, sheet, len(rows), customers)

    # שמירה חסינה: אם הקובץ פתוח באקסל (PermissionError) - לא קורסים
    # ולא מאבדים את העבודה. שומרים לקובץ חלופי עם חותמת זמן ומודיעים.
    try:
        workbook.save(output_file)
        print(f"\n📊 קובץ התוצאות נשמר: {output_file}")
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        alt_file = output_file.replace(".xlsx", f"_{timestamp}.xlsx")
        workbook.save(alt_file)
        print(f"\n⚠️  הקובץ '{output_file}' היה פתוח (אולי באקסל), אז לא ניתן היה לדרוס אותו.")
        print(f"📊 שמרתי במקום זאת לקובץ חדש: {alt_file}")
        print("   (טיפ: סגור את הקובץ באקסל לפני הרצה כדי שיישמר בשם הרגיל.)")


def _add_customer_dropdown(workbook, sheet, num_rows: int, customers):
    """
    מוסיף שתי דרכים להתאמה ידנית של לקוח:

    1. רשימה נפתחת (data validation) בעמודת "לקוח מותאם ברשימה" - לבחירה מהירה.
    2. גיליון "לקוחות" *גלוי* עם מסנן (AutoFilter) - שם יש תיבת חיפוש לפי שם
       (עובד גם ב-Office 2016/2019/2021). מחפשים שם, קוראים את המספר, ומקלידים
       אותו בעמודת "מספר לקוח במערכת" - ההפקה תזהה את הלקוח לפי המספר.

    הרשימה ממוינת לפי א״ב. הגיליון מכיל שם / מספר לקוח / ח.פ.
    """
    if not customers:
        return

    sorted_customers = sorted(customers, key=lambda c: (c.name or "").strip())

    # גיליון "לקוחות" גלוי, עם כותרות ומסנן (AutoFilter) שכולל תיבת חיפוש.
    cust_sheet = workbook.create_sheet(title=CUSTOMER_SHEET_TITLE)
    cust_sheet.sheet_view.rightToLeft = True
    cust_sheet.append(["שם לקוח", "מספר לקוח", "ח.פ"])
    for cell in cust_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for c in sorted_customers:
        cust_sheet.append([c.name, c.customer_id, c.company_id or ""])
    last = len(sorted_customers) + 1  # כולל שורת הכותרת
    cust_sheet.auto_filter.ref = f"A1:C{last}"
    cust_sheet.column_dimensions["A"].width = 32
    cust_sheet.column_dimensions["B"].width = 14
    cust_sheet.column_dimensions["C"].width = 14

    # טווח בעל-שם לשמות הלקוחות (בלי הכותרת) - מקור הרשימה הנפתחת.
    names_ref = f"'{CUSTOMER_SHEET_TITLE}'!$A$2:$A${last}"
    workbook.defined_names[CUSTOMER_NAMES_RANGE] = DefinedName(
        CUSTOMER_NAMES_RANGE, attr_text=names_ref
    )

    # מחילים את הרשימה הנפתחת על תאי עמודת "לקוח מותאם ברשימה" (שורות הנתונים).
    col_letter = openpyxl.utils.get_column_letter(
        HEADERS.index("לקוח מותאם ברשימה") + 1
    )
    last_row = num_rows + 1  # +1 בגלל שורת הכותרת בגיליון התוצאות
    dv = DataValidation(
        type="list",
        formula1=CUSTOMER_NAMES_RANGE,  # שם הטווח, בלי '=' (אחרת Excel בולע את הרשימה)
        allow_blank=True,
        showDropDown=False,  # False = להציג את חץ הרשימה (סמנטיקה הפוכה ב-OOXML)
    )
    dv.showInputMessage = True
    dv.promptTitle = "לקוח מותאם"
    dv.prompt = ("בחר מהרשימה, או חפש בגיליון 'לקוחות' (מסנן) והקלד את "
                 "מספר הלקוח בעמודת 'מספר לקוח במערכת'.")
    dv.errorTitle = "ערך לא ברשימה"
    dv.error = "בחר לקוח קיים מהרשימה, או השאר ריק והקלד מספר לקוח."
    dv.showErrorMessage = True
    sheet.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{last_row}")
