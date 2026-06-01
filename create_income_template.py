# -*- coding: utf-8 -*-
"""
==================================================================
  יוצר תבנית ריקה לקובץ תנועות הכנסה (אופציונלי)
==================================================================
קובץ תנועות ההכנסה משמש לחיזוק ההתאמה: אם יש תנועת הכנסה בסכום
זהה בדיוק לשיק והשם תואם - זה מחזק את הזיהוי של הלקוח.

הרץ אותי כדי לקבל קובץ income.xlsx מוכן עם העמודות הנכונות.
הקובץ הזה אופציונלי - הכלי עובד גם בלעדיו.

הרצה:   python create_income_template.py
==================================================================
"""

import os

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

import config


def create_template():
    if os.path.exists(config.INCOME_FILE):
        print(f"⚠️  הקובץ '{config.INCOME_FILE}' כבר קיים - לא יצרתי חדש.")
        return

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "הכנסות"
    sheet.sheet_view.rightToLeft = True

    headers = [
        config.INCOME_AMOUNT_COLUMN,
        config.INCOME_NAME_COLUMN,
        config.INCOME_DATE_COLUMN,  # "תאריך" - אופציונלי אך משפר דיוק
    ]
    fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    # שורות דוגמה - מחק ומלא את תנועות ההכנסה האמיתיות.
    # התאריך בפורמט DD/MM/YYYY (תאריך ההכנסה, בדרך כלל מוקדם לתאריך השיק).
    sheet.append([1696, "ישראל ישראלי בע\"מ", "20/05/2026"])
    sheet.append([3850, "חברת אבן וסיד בע\"מ", "22/05/2026"])

    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 14

    workbook.save(config.INCOME_FILE)
    print(f"✅ נוצרה תבנית: '{config.INCOME_FILE}'")
    print("   פתח אותה, מחק את שורות הדוגמה, ומלא את תנועות ההכנסה.")
    print("   (קובץ זה אופציונלי - הוא רק משפר את דיוק ההתאמה.)")


if __name__ == "__main__":
    create_template()
