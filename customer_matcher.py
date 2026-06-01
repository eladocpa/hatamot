# -*- coding: utf-8 -*-
"""
==================================================================
  התאמת לקוחות (ה"שדכן" של הכלי)
==================================================================
המודול הזה טוען את רשימת הלקוחות שלך מקובץ האקסל,
ולכל שיק - מנסה למצוא את הלקוח המתאים ביותר, לפי שלושה אותות:

  1. ח.פ - אם זוהה ח.פ על השיק והוא קיים ברשימה -> התאמה ודאית (החזק ביותר).
  2. שם - השוואה מטושטשת של שם המושך לרשימת הלקוחות.
  3. סכום - אם יש תנועת הכנסה בסכום זהה בדיוק -> מחזק את ההתאמה.

הכלל החשוב ביותר נשמר: לא לנחש!
אם אין התאמה ברורה, או שיש כמה לקוחות דומים -
מחזירים "דורש בדיקה ידנית" ונותנים לך להחליט.
==================================================================
"""

import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Set

import openpyxl
from thefuzz import fuzz, process

import config


@dataclass
class Customer:
    """לקוח בודד מהרשימה שלך."""
    name: str
    customer_id: str
    company_id: Optional[str] = None  # ח.פ / עוסק מורשה (אופציונלי)


@dataclass
class MatchResult:
    """תוצאת ההתאמה של שיק ללקוח."""
    matched_name: Optional[str]      # שם הלקוח שנמצא ברשימה (או None)
    matched_id: Optional[str]        # מזהה הלקוח שנמצא (או None)
    confidence: int                  # רמת ודאות 0-100
    needs_review: bool               # האם השורה דורשת בדיקה ידנית?
    reason: str                      # הסבר קצר (למה ודאי / למה לבדיקה)
    evidence: List[str] = field(default_factory=list)  # אילו אותות תמכו בהתאמה


def _find_column(header_row, *names) -> Optional[int]:
    """מוצא אינדקס עמודה לפי אחד מהשמות האפשריים. מחזיר None אם לא נמצא."""
    for index, cell_value in enumerate(header_row):
        if cell_value is not None and str(cell_value).strip() in names:
            return index
    return None


def load_customers(file_path: str = config.CUSTOMERS_FILE) -> List[Customer]:
    """
    טוען את רשימת הלקוחות מקובץ האקסל.
    מצפה לעמודות 'שם' ו'מזהה'. עמודת ח.פ היא אופציונלית.
    """
    workbook = openpyxl.load_workbook(file_path, read_only=True)
    sheet = workbook.active

    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    name_col = _find_column(header_row, config.CUSTOMERS_NAME_COLUMN)
    id_col = _find_column(header_row, config.CUSTOMERS_ID_COLUMN)
    # עמודת ח.פ - אופציונלית (גם בכמה כתיבים אפשריים).
    company_col = _find_column(
        header_row, config.CUSTOMERS_COMPANY_ID_COLUMN, "ח.פ", "חפ", "ח״פ", "עוסק מורשה"
    )

    if name_col is None or id_col is None:
        raise ValueError(
            f"לא נמצאו העמודות '{config.CUSTOMERS_NAME_COLUMN}' ו-"
            f"'{config.CUSTOMERS_ID_COLUMN}' בקובץ הלקוחות. "
            f"בדוק שהכותרות בשורה הראשונה כתובות בדיוק כך."
        )

    customers: List[Customer] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = row[name_col]
        if name is None or str(name).strip() == "":
            continue  # מדלגים על שורות ריקות
        customer_id = row[id_col]
        company_id = None
        if company_col is not None and company_col < len(row):
            raw = row[company_col]
            if raw is not None:
                company_id = _normalize_id(str(raw))
        customers.append(Customer(
            name=str(name).strip(),
            customer_id=str(customer_id).strip() if customer_id is not None else "",
            company_id=company_id,
        ))

    workbook.close()
    return customers


def load_income_index(file_path: str = config.INCOME_FILE) -> Dict[str, Set[str]]:
    """
    טוען את קובץ תנועות ההכנסה (אם קיים) ובונה מילון:
    סכום_מנורמל -> קבוצת שמות לקוחות שהיו להם הכנסה בסכום הזה.
    אם הקובץ לא קיים - מחזיר מילון ריק (פשוט מדלגים על השלב).
    """
    if not os.path.exists(file_path):
        return {}

    workbook = openpyxl.load_workbook(file_path, read_only=True)
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    amount_col = _find_column(header_row, config.INCOME_AMOUNT_COLUMN, "סכום")
    name_col = _find_column(header_row, config.INCOME_NAME_COLUMN, "שם", "לקוח")

    index: Dict[str, Set[str]] = {}
    if amount_col is None:
        workbook.close()
        return index

    for row in sheet.iter_rows(min_row=2, values_only=True):
        amount = _normalize_amount(row[amount_col]) if amount_col < len(row) else None
        if amount is None:
            continue
        name = ""
        if name_col is not None and name_col < len(row) and row[name_col] is not None:
            name = str(row[name_col]).strip()
        index.setdefault(amount, set()).add(name)

    workbook.close()
    return index


def _normalize(text: str) -> str:
    """מנקה שם לפני השוואה: מוריד תווים שלא משנים ורווחים כפולים."""
    text = text.strip()
    for ch in ['"', "'", "״", "׳", ".", ",", "-"]:
        text = text.replace(ch, " ")
    return " ".join(text.split())


def _normalize_id(raw: str) -> Optional[str]:
    """משאיר רק ספרות ממספר ח.פ (מוריד מקפים/רווחים). מחזיר None אם אין ספרות."""
    digits = re.sub(r"\D", "", str(raw))
    return digits or None


def _normalize_amount(raw) -> Optional[str]:
    """
    מנרמל סכום לצורת השוואה אחידה: מוריד ₪, פסיקים ורווחים,
    ולוקח ערך מוחלט (כי הכנסה חיובית מול שיק - מתאימים לפי גודל).
    מחזיר מחרוזת כמו '1503.66', או None אם אי אפשר לפענח.
    """
    if raw is None:
        return None
    text = str(raw).replace("₪", "").replace(",", "").replace(" ", "").strip()
    text = text.replace("‏", "").replace("‎", "")  # סימני כיווניות
    if text == "":
        return None
    try:
        value = abs(float(text))
    except ValueError:
        return None
    # מנרמלים לשתי ספרות אחרי הנקודה, בלי אפסים מיותרים בסוף.
    return f"{value:.2f}".rstrip("0").rstrip(".")


class CustomerMatcher:
    """מתאים שיקים ללקוחות לפי ח.פ, שם, וסכום הכנסה."""

    def __init__(self, customers: List[Customer], income_index: Optional[Dict] = None):
        self.customers = customers
        self.income_index = income_index or {}
        # מילון שם מנורמל -> לקוח
        self._normalized_map = {_normalize(c.name): c for c in customers}
        self._normalized_names = list(self._normalized_map.keys())
        # מילון ח.פ -> לקוח (רק ללקוחות שיש להם ח.פ)
        self._company_id_map = {
            c.company_id: c for c in customers if c.company_id
        }

    def match(
        self,
        detected_name: Optional[str],
        company_id: Optional[str] = None,
        amount: Optional[str] = None,
    ) -> MatchResult:
        """
        מתאים שיק ללקוח. משלב שלושה אותות: ח.פ, שם, וסכום הכנסה.
        מחיל את הכלל: בספק - לבדיקה ידנית.
        """
        evidence: List[str] = []

        # ---- אות 1: ח.פ (החזק ביותר) ----
        # אם זוהה ח.פ על השיק והוא קיים ברשימת הלקוחות -> התאמה ודאית.
        norm_company = _normalize_id(company_id) if company_id else None
        if norm_company and norm_company in self._company_id_map:
            cust = self._company_id_map[norm_company]
            evidence.append(f"ח.פ תואם ({norm_company})")
            return MatchResult(
                matched_name=cust.name, matched_id=cust.customer_id,
                confidence=100, needs_review=False,
                reason="התאמה ודאית לפי ח.פ", evidence=evidence,
            )

        # ---- אות 2: שם (השוואה מטושטשת) ----
        if not detected_name or detected_name.strip() == "":
            return MatchResult(
                matched_name=None, matched_id=None, confidence=0,
                needs_review=True, reason="לא זוהה שם על השיק", evidence=evidence,
            )

        normalized_detected = _normalize(detected_name)
        top_matches = process.extract(
            normalized_detected, self._normalized_names,
            scorer=fuzz.token_sort_ratio, limit=2,
        )
        if not top_matches:
            return MatchResult(
                matched_name=None, matched_id=None, confidence=0,
                needs_review=True, reason="אין לקוחות ברשימה", evidence=evidence,
            )

        best_name, best_score = top_matches[0]
        best_customer = self._normalized_map[best_name]
        second_score = top_matches[1][1] if len(top_matches) > 1 else 0
        gap = best_score - second_score

        # ---- אות 3: סכום הכנסה (מחזק) ----
        # אם יש תנועת הכנסה בסכום זהה, ושם הלקוח שלה תואם למועמד -
        # זה מחזק את ההתאמה (מוסיף נקודות ביטחון).
        income_boost = 0
        norm_amount = _normalize_amount(amount) if amount else None
        if norm_amount and norm_amount in self.income_index:
            income_names = self.income_index[norm_amount]
            # בודקים אם שם הלקוח המועמד מופיע בין ההכנסות באותו סכום.
            for income_name in income_names:
                if income_name and fuzz.token_sort_ratio(
                    _normalize(income_name), best_name
                ) >= config.MATCH_CONFIDENCE_THRESHOLD:
                    income_boost = 8
                    evidence.append(f"סכום זהה בתנועות הכנסה ({norm_amount})")
                    break
            else:
                # סכום קיים אך השם לא תואם - אות חלש, רק לתיעוד.
                if income_names:
                    evidence.append(f"סכום {norm_amount} קיים בהכנסות (שם אחר)")

        effective_score = min(100, best_score + income_boost)
        if best_score >= config.MATCH_CONFIDENCE_THRESHOLD or income_boost:
            evidence.insert(0, f"שם דומה ({best_score}%)")

        # ---- הכרעה ----
        # התאמה חלשה מדי גם אחרי החיזוק -> בדיקה ידנית.
        if effective_score < config.MATCH_CONFIDENCE_THRESHOLD:
            return MatchResult(
                matched_name=best_customer.name, matched_id=best_customer.customer_id,
                confidence=effective_score, needs_review=True,
                reason=f"התאמה חלשה ({best_score}%)", evidence=evidence,
            )

        # שני לקוחות דומים מדי, ואין אות מכריע (ח.פ/סכום) -> בדיקה ידנית.
        if (gap < config.MATCH_AMBIGUITY_GAP
                and second_score >= config.MATCH_CONFIDENCE_THRESHOLD
                and income_boost == 0):
            return MatchResult(
                matched_name=best_customer.name, matched_id=best_customer.customer_id,
                confidence=effective_score, needs_review=True,
                reason=f"כמה לקוחות דומים (פער {gap}% בלבד)", evidence=evidence,
            )

        # התאמה ודאית!
        reason = "התאמה ודאית"
        if income_boost:
            reason = "התאמה ודאית (שם + סכום הכנסה)"
        return MatchResult(
            matched_name=best_customer.name, matched_id=best_customer.customer_id,
            confidence=effective_score, needs_review=False,
            reason=reason, evidence=evidence,
        )
